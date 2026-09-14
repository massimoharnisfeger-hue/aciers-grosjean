"""
Classeur de collecte des données pour les rendus 3D des fiches produits.

Lit lib/catalogue.ts et produit _DEPOT/rendus-3d/collecte-donnees-produits.xlsx :
une feuille par vague de production, une ligne par produit, pré-remplie avec
les cotes du catalogue, la fiche technique source et une finition proposée.
Le propriétaire n'a qu'à confirmer ou corriger les cases jaunes.

Usage : python scripts/generer-collecte-rendus.py
"""

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOGUE = os.path.join(RACINE, "lib", "catalogue.ts")
SORTIE = os.path.join(RACINE, "_DEPOT", "rendus-3d", "collecte-donnees-produits.xlsx")

ENCRE, JAUNE, BRUME, NUAGE = "333642", "FFD500", "D1D6DA", "F4F6F7"

FINITIONS = [
    "Noir laminé à chaud (calamine)",
    "Décapé-huilé",
    "Gris clair huilé (laminé à froid)",
    "Brut étiré à froid (clair)",
    "Galvanisé à chaud",
    "Galvanisé sendzimir (sans fleurs)",
    "Galvanisé à fleurs",
    "Brut crénelé (légère rouille)",
    "Aluminium brut",
    "Aluminium avec film de protection",
    "Inox brossé grain 320",
    "Inox brossé + film de protection",
    "Corten neuf (gris-bleu)",
    "Corten pré-patiné (rouille)",
    "Thermolaqué / prélaqué RAL",
    "Parement imitation bois",
    "Photo fournisseur (pas de 3D)",
    "Autre (préciser en remarque)",
]

# catégorie -> (vague, finition proposée, fiche technique source)
REGLES = [
    ("/acier/poutrelles/ipe", 1, "Noir laminé à chaud (calamine)", "POUTRELLE IPE ACIERS GROSJEAN 4(1).pdf"),
    ("/acier/poutrelles/hea", 1, "Noir laminé à chaud (calamine)", "INFOSTEEL - GRYMAFER - VM2013 - HEA.pdf"),
    ("/acier/poutrelles/heb", 1, "Noir laminé à chaud (calamine)", "INFOSTEEL - GRYMAFER - VM2013 - HEB.pdf"),
    ("/acier/poutrelles/upn", 1, "Noir laminé à chaud (calamine)", "INFOSTEEL - GRYMAFER - VM2013 - UPN.pdf"),
    ("/acier/profiles/corniere-egale", 1, "Noir laminé à chaud (calamine)", "CE EGALE LAC.pdf"),
    ("/acier/profiles/corniere-inegale", 1, "Noir laminé à chaud (calamine)", "CI-3.pdf"),
    ("/acier/profiles/fer-t", 1, "Noir laminé à chaud (calamine)", "FER-T ACIERS GROSJEAN 4(1).pdf"),
    ("/acier/profiles/large-plat", 1, "Noir laminé à chaud (calamine)", "LARGE PLAT 4.pdf"),
    ("/acier/profiles/plat", 1, "Noir laminé à chaud (calamine)", "3(2).pdf"),
    ("/acier/profiles/rond-plein", 1, "Noir laminé à chaud (calamine)", "ROND LISSE 3(1).pdf"),
    ("/acier/profiles/carre-plein", 1, "Noir laminé à chaud (calamine)", "CARRE PLEIN.pdf"),
    ("/acier/tubes/tube-carre", 1, "Noir laminé à chaud (calamine)", "TUBE CARRE ACIERS GROSJEAN 4.pdf"),
    ("/acier/tubes/tube-rectangulaire", 1, "Noir laminé à chaud (calamine)", "TUBE CARRE ACIERS GROSJEAN 4.pdf (à vérifier)"),
    ("/acier/tubes/tube-rond", 1, "Noir laminé à chaud (calamine)", "TUBE ROND ACIERS GROSJEAN 4.pdf"),
    ("/acier/toles/tole-laminee-a-chaud", 1, "Noir laminé à chaud (calamine)", "ACIERS GROSJEAN - TOLES LAC - 3.pdf"),
    ("/acier/armatures-beton/rond-a-beton", 1, "Brut crénelé (légère rouille)", ""),
    ("/acier/armatures-beton/treillis", 1, "Brut crénelé (légère rouille)", "TREILLIS SOUDE EN ACIER POUR LA CONSTRUCTION - ACIERS GROSJEAN-3.pdf + AG - TS - *.pdf"),
    ("/acier/toles/tole-laminee-a-froid", 2, "Gris clair huilé (laminé à froid)", ""),
    ("/acier/toles/tole-galvanisee", 2, "Galvanisé sendzimir (sans fleurs)", ""),
    ("/acier/toles/tole-corten", 2, "Corten neuf (gris-bleu)", ""),
    ("/acier/toles/tole-quarto", 2, "Noir laminé à chaud (calamine)", ""),
    ("/acier/toles/tole-larmee", 2, "Noir laminé à chaud (calamine)", ""),
    ("/acier/toles/tole-perforee", 2, "Noir laminé à chaud (calamine)", ""),
    ("/aluminium/", 2, "Aluminium brut", ""),
    ("/inox/", 2, "Inox brossé grain 320", ""),
    ("/jardin-cloture/clotures", 3, "Thermolaqué / prélaqué RAL", "GAG - FT-CG64-Plis205.pdf / GAG - FT-CP40-Plis205.pdf + notices de pose"),
    ("/jardin-cloture/amenagement/bordures", 3, "", ""),
    ("/quincaillerie/caillebotis-marches", 3, "Galvanisé à chaud", "FT-PLANCHER SECURIT PCP.pdf / FT-Marches D'Escaliers ACHIL O2-Acier 240 PCP.pdf"),
    ("/toiture-bardage/toles-profilees", 3, "Thermolaqué / prélaqué RAL", "ag-tole-profilee-30-200.pdf / ACIERS GROSJEAN - TP - 30-200.pdf"),
    ("/toiture-bardage/panneaux-isoles", 3, "Thermolaqué / prélaqué RAL", "lattonedil-eurocopre-monolamiera-ag.pdf"),
    ("/toiture-bardage/bardage", 3, "Parement imitation bois", "ag-tasseau-40x40-maxi-imitation-bois.pdf"),
    ("/quincaillerie/", 4, "Photo fournisseur (pas de 3D)", ""),
]

VAGUES = {
    1: ("Vague 1 - acier", "Profilés, tubes, tôles noires, armatures. Données complètes : production dès confirmation de la finition."),
    2: ("Vague 2 - tôles, alu, inox", "Cotes présentes ; finition et motif à confirmer."),
    3: ("Vague 3 - clôture, toiture", "Modélisation plus lourde, d'après les fiches techniques fournisseurs."),
    4: ("Vague 4 - produits de marque", "Pas de 3D : photos des fournisseurs."),
}


def regle(categorie):
    for prefixe, vague, finition, source in REGLES:
        if categorie.startswith(prefixe):
            return vague, finition, source
    return 4, "", ""


def anomalies(nom, categorie, specs):
    notes = []
    if categorie == "/aluminium/profiles/profil-u":
        ep = dict(specs).get("Épaisseur", "")
        m = re.search(r"x(\d+(?:,\d+)?)mm", nom)
        if m and ep and ep.split()[0] != m.group(1):
            notes.append(f"Le site affiche une épaisseur de {ep}, le nom indique {m.group(1)} mm : à corriger.")
    if categorie.startswith(("/aluminium/toles", "/inox/toles")) and "oxycoupage" in dict(specs).get("Découpe", ""):
        notes.append("Le site propose l'oxycoupage, impossible sur ce métal : à corriger.")
    if "VERT RAL 7016" in nom:
        notes.append("RAL 7016 est un gris anthracite, pas un vert : à vérifier.")
    if nom.startswith("PANNEAUX MEDIUM 3D"):
        notes.append("Aucune fiche technique pour le modèle MEDIUM 3D : à fournir.")
    if "aléatoire" in nom:
        notes.append("Motif aléatoire : photo ou plan nécessaire.")
    if categorie == "/jardin-cloture/clotures/fixations":
        notes.append("Photo ou plan coté nécessaire.")
    if categorie == "/jardin-cloture/amenagement/bordures":
        notes.append("Préciser la forme (hauteur 150, épaisseur ?, retour 25 ?) et les piquets.")
    return " ".join(notes)


def lire_catalogue():
    produits = []
    with open(CATALOGUE, encoding="utf-8") as f:
        for ligne in f:
            m = re.match(r'\s*"([^"/]+)": \{ slug: "[^"]+", nom: "(.*?)", categorie: "(.*?)"', ligne)
            if not m:
                continue
            specs = re.findall(r'\{ label: "(.*?)", valeur: "(.*?)" \}', ligne)
            produits.append({"slug": m.group(1), "nom": m.group(2), "categorie": m.group(3), "specs": specs})
    return produits


def entete(ws, colonnes):
    fin = Side(style="thin", color=BRUME)
    for i, (titre, largeur, a_remplir) in enumerate(colonnes, start=1):
        c = ws.cell(row=1, column=i, value=titre)
        c.font = Font(bold=True, color=ENCRE if a_remplir else "FFFFFF")
        c.fill = PatternFill("solid", fgColor=JAUNE if a_remplir else ENCRE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=fin)
        ws.column_dimensions[c.column_letter].width = largeur
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"


def main():
    produits = lire_catalogue()
    wb = Workbook()

    # ---- mode d'emploi
    mode = wb.active
    mode.title = "Mode d'emploi"
    lignes = [
        ("Collecte des données pour les photos 3D des fiches produits", True),
        ("", False),
        ("Colonnes grises : déjà remplies depuis le site. Colonnes jaunes : à confirmer ou corriger.", False),
        ("« Finition proposée » est une supposition : choisissez la finition réelle en stock dans la liste.", False),
        ("« Cotes exactes ? » : oui si les données actuelles correspondent au stock, sinon non + correction.", False),
        ("Une case « Remarque » déjà remplie signale une erreur ou un manque repéré : merci d'y répondre.", False),
        ("", False),
        ("Ordre de production", True),
    ] + [(f"{titre} : {texte}", False) for titre, texte in VAGUES.values()] + [
        ("", False),
        ("Aussi à fournir (voir _DOCS/BRIEF-RENDUS-3D.md) : photos de référence des finitions, feuille « Photos matières ».", False),
    ]
    for i, (texte, gras) in enumerate(lignes, start=1):
        c = mode.cell(row=i, column=1, value=texte)
        c.font = Font(bold=gras, size=14 if i == 1 else 11, color=ENCRE)
    mode.column_dimensions["A"].width = 120

    # ---- listes (feuille masquée pour les menus déroulants)
    listes = wb.create_sheet("listes")
    for i, f in enumerate(FINITIONS, start=1):
        listes.cell(row=i, column=1, value=f)
    listes.sheet_state = "hidden"

    colonnes = [
        ("Référence (slug)", 34, False),
        ("Produit", 46, False),
        ("Catégorie", 30, False),
        ("Données actuelles du site", 60, False),
        ("Fiche technique source", 34, False),
        ("Finition proposée", 28, False),
        ("Finition réelle en stock", 30, True),
        ("Cotes exactes ?", 11, True),
        ("Correction des cotes", 26, True),
        ("Détails visibles (soudure, film, motif, couleur…)", 34, True),
        ("Remarque", 46, True),
    ]
    fond_pair = PatternFill("solid", fgColor=NUAGE)

    for vague, (titre, _) in VAGUES.items():
        ws = wb.create_sheet(titre[:31])
        entete(ws, colonnes)
        dv_fin = DataValidation(type="list", formula1=f"=listes!$A$1:$A${len(FINITIONS)}", allow_blank=True)
        dv_oui = DataValidation(type="list", formula1='"oui,non"', allow_blank=True)
        ws.add_data_validation(dv_fin)
        ws.add_data_validation(dv_oui)
        rang = 2
        for p in produits:
            v, finition, source = regle(p["categorie"])
            if v != vague:
                continue
            valeurs = [
                p["slug"], p["nom"], p["categorie"],
                " · ".join(f"{k} : {val}" for k, val in p["specs"]),
                source, finition, "", "", "", "", anomalies(p["nom"], p["categorie"], p["specs"]),
            ]
            for col, val in enumerate(valeurs, start=1):
                c = ws.cell(row=rang, column=col, value=val)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if rang % 2 == 0 and not colonnes[col - 1][2]:
                    c.fill = fond_pair
                if col == 11 and val:
                    c.font = Font(color="B45309", bold=True)
            dv_fin.add(f"G{rang}")
            dv_oui.add(f"H{rang}")
            rang += 1
        ws.auto_filter.ref = f"A1:K{rang - 1}"

    # ---- photos de référence des finitions
    ph = wb.create_sheet("Photos matières")
    entete(ph, [("Finition", 34, False), ("Produit à photographier", 44, False),
                ("Nom du fichier", 40, False), ("Fait ?", 10, True), ("Remarque", 40, True)])
    photos = [
        ("Noir laminé à chaud (calamine)", "Poutrelle IPE ou HEA, avec une extrémité sciée", "calamine"),
        ("Tube acier noir", "Tube carré, vue de l'extrémité (cordon de soudure intérieur)", "tube-noir"),
        ("Plat / rond / carré plein", "Barres en rack, vue d'ensemble + gros plan", "barres-pleines"),
        ("Tôle laminée à froid", "Coin de tôle, lumière rasante", "tole-laf"),
        ("Galvanisé", "Tôle galvanisée : surface en gros plan (fleurs ou non)", "galvanise"),
        ("Corten", "Tôle corten telle que livrée", "corten"),
        ("Tôle larmée", "Motif en gros plan, pièce de monnaie posée pour l'échelle", "tole-larmee"),
        ("Tôle perforée", "Perforation R5 T8 et C10 U15, gros plan", "tole-perforee"),
        ("Aluminium", "Tube ou cornière alu, avec film éventuel", "aluminium"),
        ("Tôle striée alu", "Motif en gros plan", "tole-striee-alu"),
        ("Inox brossé", "Tube ou tôle inox, sens du brossage visible", "inox-brosse"),
        ("Treillis / rond à béton", "Panneau ou barres en stock", "treillis"),
        ("Caillebotis", "Caillebotis galvanisé, gros plan de la maille", "caillebotis"),
        ("Clôture", "Panneau + poteau montés, avec la fixation", "cloture"),
    ]
    for i, (fin, produit, nom) in enumerate(photos, start=2):
        for col, val in enumerate([fin, produit, f"_DEPOT/images/references-matieres/{nom}-1.jpg"], start=1):
            ph.cell(row=i, column=col, value=val).alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
    wb.save(SORTIE)
    total = {v: sum(1 for p in produits if regle(p["categorie"])[0] == v) for v in VAGUES}
    nb_anomalies = sum(1 for p in produits if anomalies(p["nom"], p["categorie"], p["specs"]))
    print(f"{len(produits)} produits -> {SORTIE}")
    print("par vague :", total, "| lignes avec remarque :", nb_anomalies)


main()
